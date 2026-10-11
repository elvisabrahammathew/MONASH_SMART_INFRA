"""Build the MMA3001 submission draft from reproducible project results."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report" / "MMA3001_Laboratory_Availability_Report.docx"

def table(d, headers, rows):
    t=d.add_table(rows=1, cols=len(headers)); t.style="Light Shading Accent 1"
    for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row): c.text=str(v)

def para(d, value): d.add_paragraph(value)

def main():
    d=Document(); s=d.sections[0]; s.top_margin=s.bottom_margin=Inches(.7)
    d.styles["Normal"].font.name="Aptos"; d.styles["Normal"].font.size=Pt(10.5)
    h=d.add_heading("Smart laboratory availability forecasting for Building 60, G.25",0); h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph("MMA3001 Numerical Methods and Machine Learning Project\nStudent: Elvis Abraham Mathew\nOctober 2026"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading("Abstract",1); para(d,"This project forecasts whether Building 60, room G.25 (Keenan Lab) will be available 30 minutes ahead for ad hoc laboratory booking, study-space allocation and room-readiness decisions. A sensor-derived headcount forecast is converted to available when it is at or below 0.5 people. Irregular events were aggregated to 15-minute intervals. Ridge regression was selected on validation, while persistence had the lowest final-test MAE (0.461 people). The selected ridge availability signal achieved 64.6% accuracy and 88.1% precision, meaning an available indication was usually correct but some available periods were missed. The repository includes reproducible code, tests, documentation and an AI-use disclosure.")
    d.add_heading("1. Engineering problem and scope",1); para(d,"The question is: can G.25 availability 30 minutes ahead be inferred from headcount forecasts more reliably than simple operational baselines? The input is timestamped sensor-derived headcount and the output is an available/not-available signal. A count at or below 0.5 people means available. The intended use is ad hoc laboratory booking, study-space allocation and room-readiness decisions; it is not a confirmed booking, access-control decision, or independent sensor validation. Environmental variables are excluded because the supplied metadata does not validate a co-located sensor for G.25.")
    d.add_heading("2. Data audit and preparation",1); para(d,"The source archive contains 643,290 G.25 events between 1 May and 31 December 2024. Raw headcount ranges from 0–35. Events were resampled to 15-minute bins; the last state was carried only across gaps of up to 60 minutes. Longer gaps remain missing, so absent data are never interpreted as an empty room.")
    table(d,["Audit quantity","Value"],[["15-minute bins","23,516"],["Eligible 30-minute forecasts","19,838"],["Median event gap","4 seconds"],["Longest gap","77,778 s (21.6 h)"],["Missing after carry rule","3,146 bins"],["Forward-filled under rule","3,780 bins"]])
    d.add_heading("3. Computational method",1); para(d,"The pipeline reads the archive in chunks, filters the unique G.25 floor-space identifier, and creates features available at issue time: current headcount, lags of 1, 2, 4 and 8 intervals, rolling statistics, time-of-day sine/cosine, weekday and weekend. The target is headcount two 15-minute intervals later. A negative shift constructs only the label and is excluded from predictors, preventing future-data leakage.")
    para(d,"A 60/20/20 chronological split preserves time order. Ridge regression, decision-tree regression and SVR were compared with persistence and weekday/time baselines. MAE measures typical absolute count error, RMSE weights large errors, signed bias detects systematic error, and runtime represents operational cost.")
    d.add_heading("4. Validation and model selection",1); para(d,"Validation, never the final test, selected the method. At 30 minutes ridge had the lowest validation MAE (1.080 people), ahead of SVR (1.131), decision tree (1.140) and persistence (1.152). Ridge was selected for accuracy, interpretability and roughly 0.04-second fitting; SVR required about 13 seconds. Five automated tests verify resampling, future-target alignment, chronological splitting, hand-calculated metrics and a synthetic end-to-end workflow.")
    d.add_heading("5. Final results and sensitivity",1); table(d,["Horizon","Validation winner","Final-test best MAE","MAE","Interpretation"],[["15 min","Ridge","Persistence","0.407","Lowest typical error"],["30 min","Ridge","Persistence","0.461","Main forecast horizon"],["60 min","Ridge","Persistence","0.547","Error rises with horizon"]])
    para(d,"For the main 30-minute task, final-test MAE was 0.461 for persistence, 0.507 for SVR, 0.579 for ridge, 0.591 for the decision tree and 1.168 for the weekday/time baseline. SVR had the lowest RMSE (0.794 people), compared with 0.837 for ridge and 0.944 for persistence. Ridge overpredicted by 0.272 people. The final period was less variable than development periods, so persistence won MAE despite ridge being correctly selected from validation. The model was not reselected after test evaluation.")
    d.add_heading("6. Engineering interpretation and limitations",1); para(d,"A 15-minute persistence forecast is the most accurate tested practical baseline; uncertainty rises at 30 and 60 minutes. The ridge availability signal has 88.1% precision but 53.8% recall, so it is useful as a cautious advisory indicator rather than a booking guarantee. Limitations are the one-room scope, no independent ground truth, a 21.6-hour outage, no booking-system integration, possible metadata-mapping issues, and no validated environmental co-location. Future work should compare against bookings and manual counts, analyse transition errors and evaluate more rooms.")
    d.add_heading("7. Reproducibility, ethics and AI use",1); para(d,"Raw data are excluded from version control. The repository provides a runbook, data dictionary, HTML documentation, test report and commands to regenerate outputs. The source is treated as sensor-derived and potentially restricted. Codex assisted planning, code drafting and documentation, while project evidence and claims were checked against the outputs. The full disclosure is in docs/AI_USAGE.md.")
    d.add_heading("References",1); para(d,"MMA3001 (2026), Project Brief and Dataset 2 project information, Monash University.\nScikit-learn developers (2026), scikit-learn user guide. https://scikit-learn.org/stable/ (accessed 9 October 2026).")
    d.add_heading("Appendix: repository evidence",1); para(d,"Main entry: src/occupancy/pipeline.py. Sensitivity: python -m occupancy.sensitivity. Test report: results/test-report.html. API documentation: docs/html/index.html. Horizon metrics: results/horizon_sensitivity.csv.")
    OUT.parent.mkdir(exist_ok=True); d.save(OUT); print(OUT)
if __name__=="__main__": main()
